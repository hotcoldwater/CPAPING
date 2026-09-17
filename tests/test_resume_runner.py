import base64,copy,io,unittest
from email import policy
from email.parser import BytesParser
from unittest.mock import Mock,patch
from reportlab.pdfgen import canvas
from career import resume_runner as r,mail,matching

class ResumeTests(unittest.TestCase):
 def setUp(self):
  out=io.BytesIO();c=canvas.Canvas(out);c.drawString(40,750,'Synthetic uploaded resume - exact bytes');c.save();self.raw=out.getvalue()
  self.file={'id':'file','user_id':'u','kind':'resume','mime':'application/pdf','name':'resume.pdf','data_base64':base64.b64encode(self.raw).decode()}
  self.source='회사명: 예시법인\n제출서류: 이력서(자유양식)\n이메일 지원: hr@example.com\n마감일: 2099-12-31'
  self.rule={'user_id':'u','resume_file_id':'file','applicant_name':'지원자','mail_subject_template':'[입사지원] {법인} - {이름}','mail_body_template':'이력서를 첨부합니다. {이름}', 'updated_at':'r1','enabled':True,'mode':'auto','consent_version':'resume-auto-v1','enabled_since':'2026-01-01','filters':{'scope':matching.RESUME_SCOPE,'employment':['Full Time','Part Time']}}
  self.post={'id':1,'company_name':'예시법인','title':'채용','detail_url':'https://www.kicpa.or.kr/test','source':'kicpa:trainee','posting_type':'entry','audience':'cpa','is_target':True,'employment_type':'Full Time','first_seen_at':'2026-09-14'}
  self.account={'user_id':'u','email':'self@example.com','connected_at':'m1'}
  self.app={'id':'app','user_id':'u','posting_id':1,'origin':'rule','version':1,'created_at':'2026-09-13','status':'preparing','snapshot':{'flow':r.FLOW}}
  self.tables={'career_rules':self.rule,'career_files':self.file,'job_postings':self.post,'career_mail_accounts':self.account,'career_applications':self.app}
  self.db=Mock();self.db.one.side_effect=lambda table,**kw:copy.deepcopy(self.tables.get(table))
 def prepare(self):
  with patch('career.resume_runner.source',return_value=self.source),patch('career.requirements_ai.analyze',return_value={'state':'classified','subject':{'kind':'free'},'filename':{'kind':'free'},'documents':{'kind':'free'},'method':'email','recipient':'hr@example.com','file_format':'pdf','uncertainty':[],'blockers':[]}):r.prepare(self.db,self.app)
  data=self.db.update.call_args.args[1];self.app.update(data);return data
 def test_explicit_free_resume_queues_after_ai_analysis_and_attaches_exact_file(self):
  d=self.prepare();self.assertEqual(d['status'],'queued');self.assertEqual(d['document_id'],'file');self.assertTrue(d['snapshot']['auto'])
  content=mail.mime_message(self.app,self.account,self.file);m=BytesParser(policy=policy.default).parsebytes(content)
  attachments=list(m.iter_attachments());self.assertEqual(attachments[0].get_payload(decode=True),self.raw);self.assertEqual(attachments[0].get_filename(),'resume.pdf')
 def test_extra_documents_and_subject_rules_require_review(self):
  for text in ['\n자기소개서 제출','\n메일 제목: 이름_직무','\n지정 양식 첨부파일','\n참조: other@example.com','\n경력기술서 제출','\n온라인 접수','\n온라인으로만 지원','\n제목은 성명으로 작성','\n메일 본문에 희망급여 기재','\n주민등록등본 제출']:
   with self.subTest(text=text):self.assertTrue(r.requirements(self.source+text,self.file)[1])
 def test_no_free_format_or_multiple_addresses_never_auto_send(self):
  self.assertTrue(r.requirements(self.source.replace('이메일 지원:', '문의 이메일:'),self.file)[1]);self.assertTrue(r.requirements(self.source.replace('(자유양식)',''),self.file)[1]);self.assertEqual(len(r.requirements(self.source+' 문의: second@example.com',self.file)[0]),2)
 def test_invalid_resume_rejected_before_preparing(self):
  self.file['data_base64']=base64.b64encode(b'%PDF-1.4 invalid').decode()
  with self.assertRaises(ValueError):self.prepare()
 def test_latest_source_past_deadline_rejected(self):
  with self.assertRaises(ValueError):r.deadline_check('마감일: 2020-01-01')
  with self.assertRaises(ValueError):r.deadline_check('이번 채용이 종료되었습니다.')
 def test_auto_send_uses_history_and_exact_uploaded_file(self):
  self.prepare();self.app['status']='sending';self.db.reset_mock()
  with patch('career.resume_runner.source',return_value=self.source),patch('career.resume_runner.delivery.deliver',return_value=('mid','did')) as deliver:r.send_one(self.db,self.app)
  deliver.assert_called_once();self.assertEqual(deliver.call_args.kwargs['mode'],'auto');self.assertEqual(deliver.call_args.args[3][0]['data_base64'],self.file['data_base64']);self.assertEqual(self.db.update.call_args.args[1]['status'],'sent')
 def test_changed_rule_or_post_blocks_before_send(self):
  self.prepare();self.app['status']='sending'
  with patch('career.resume_runner.source',return_value=self.source+' changed'),patch('career.resume_runner.delivery.deliver') as deliver:r.send_one(self.db,self.app);deliver.assert_not_called()
  self.rule['updated_at']='r2'
  with patch('career.resume_runner.delivery.deliver') as deliver:r.send_one(self.db,self.app);deliver.assert_not_called()
 def test_unknown_delivery_does_not_queue_retry(self):
  self.prepare();self.app['status']='sending'
  with patch('career.resume_runner.source',return_value=self.source),patch('career.resume_runner.delivery.deliver',side_effect=mail.DeliveryUnknown('unknown')):r.send_one(self.db,self.app)
  self.assertEqual(self.db.update.call_args.args[1]['status'],'delivery_unknown')
 def test_manual_request_stays_review_even_with_auto_rule(self):
  self.app['origin']='manual';self.assertEqual(self.prepare()['status'],'review')
 def test_employment_filter_and_career_exclusion(self):
  self.assertTrue(matching.matches_resume(self.post,self.rule))
  self.rule['filters']['employment']=['Part Time']
  self.assertFalse(matching.matches_resume(self.post,self.rule))
  self.post['employment_type']='Part Time'
  self.assertTrue(matching.matches_resume(self.post,self.rule))
  for change in [{'source':'kicpa:cpa'},{'posting_type':'experienced'},{'posting_type':'mixed'},{'audience':'staff'},{'is_target':False},{'career_min_years':1},{'employment_type':None},{'original_id':2},{'is_expired':True},{'first_seen_at':'2025-01-01'}]:
   with self.subTest(change=change):self.assertFalse(matching.matches_resume({**self.post,**change},self.rule))
  for filters in [{'all_firms':True},{'employment':['Part Time']},{'scope':matching.RESUME_SCOPE,'employment':[]},{'scope':matching.RESUME_SCOPE,'employment':['Career']}]:
   with self.subTest(filters=filters):self.assertFalse(matching.matches_resume(self.post,{**self.rule,'filters':filters}))
 def test_history_cutoff_and_disabled_rules_never_backfill(self):
  self.assertFalse(matching.matches_resume(self.post,{**self.rule,'history_since':'2026-09-15'}))
  self.assertFalse(matching.matches_resume(self.post,{**self.rule,'enabled':False}))
  self.assertTrue(matching.matches_resume(self.post,{**self.rule,'history_since':'2026-09-13'}))
 def test_initial_failure_mail_has_real_name_and_company(self):
  value=r.initial_message(self.rule,self.post)
  self.assertEqual(value['subject'],'[입사지원] 예시법인 - 지원자')
  self.assertEqual(value['body'],'이력서를 첨부합니다. 지원자')
  self.assertEqual(value['snapshot']['company'],'예시법인')
 def test_matcher_only_reads_rules_and_postings(self):
  self.db.all.side_effect=lambda table,**kw: {'career_rules':[self.rule],'job_postings':[self.post,{**self.post,'id':2,'source':'kicpa:cpa'}]}[table]
  self.db.one.return_value=None;self.db.one.side_effect=None
  r.match_new(self.db)
  self.db.insert.assert_called_once();self.assertEqual(self.db.insert.call_args.args[1]['posting_id'],1)
 def test_career_rejected_in_preparation(self):
  self.post['source']='kicpa:cpa'
  with self.assertRaisesRegex(ValueError,'경력직'):self.prepare()
  self.db.update.assert_not_called()
 def test_reclassified_post_and_changed_employment_block_before_delivery(self):
  self.prepare();self.app['status']='sending'
  for change in [{'source':'kicpa:cpa'},{'posting_type':'experienced'},{'career_min_years':2}]:
   with self.subTest(change=change),patch.dict(self.post,change),patch('career.resume_runner.delivery.deliver') as deliver:
    r.send_one(self.db,self.app);deliver.assert_not_called();self.assertEqual(self.db.update.call_args.args[1]['status'],'blocked')
  self.rule['filters']['employment']=['Part Time']
  with patch('career.resume_runner.delivery.deliver') as deliver:r.send_one(self.db,self.app);deliver.assert_not_called()
 def test_manual_career_application_remains_reviewable(self):
  self.post['source']='kicpa:cpa';self.app['origin']='manual'
  self.assertEqual(self.prepare()['status'],'review')
 def test_disabled_worker_has_no_db_or_ai_activity(self):
  with patch.dict('os.environ',{'CAREER_RESUME_ENABLED':'false'}),patch('career.resume_runner.DB') as db:r.main();db.assert_not_called()

if __name__=='__main__':unittest.main()

class FormatSelectionTests(unittest.TestCase):
 setUp=ResumeTests.setUp
 def test_docx_selection_and_designated_names_use_exact_original_bytes(self):
  import docx
  out=io.BytesIO();d=docx.Document();d.add_paragraph('Synthetic Word support file');d.save(out)
  word={**self.file,'id':'word','name':'resume.docx','mime':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','data_base64':base64.b64encode(out.getvalue()).decode()}
  self.rule['resume_docx_file_id']='word'
  self.db.one.side_effect=lambda table,**kw:copy.deepcopy(word if table=='career_files' and kw.get('id')=='eq.word' else self.tables.get(table))
  result={'state':'classified','subject':{'kind':'designated','template':'신입_{이름}'},'filename':{'kind':'designated','template':'{이름}_지원서'},'documents':{'kind':'free'},'method':'email','recipient':'hr@example.com','file_format':'docx','uncertainty':[],'blockers':[]}
  with patch('career.resume_runner.source',return_value=self.source),patch('career.requirements_ai.analyze',return_value=result):r.prepare(self.db,self.app)
  prepared=self.db.update.call_args.args[1];self.assertEqual(prepared['document_id'],'word');self.assertEqual(prepared['subject'],'신입_지원자');self.assertEqual(prepared['snapshot']['attachments'][0]['name'],'지원자_지원서.docx');self.assertEqual(prepared['snapshot']['document_hash'],r.digest(word['data_base64']))
 def test_square_and_legacy_tokens_resolve_for_actual_delivery(self):
  for template in ('[입사지원] [회계법인] - [이름] / [공고]', '[입사지원] {법인} - {이름} / {공고}'):
   self.assertEqual(r.message(template,self.rule,self.post),'[입사지원] 예시법인 - 지원자 / 채용')
  self.assertEqual(r.message('[이름] 드림\n[회계법인] [공고]',self.rule,self.post),'지원자 드림\n예시법인 채용')
 def test_birth_year_from_registered_member_or_block(self):
  self.assertEqual(r.message('{이름}_{출생년도}',{**self.rule,'member':{'birth_date':'1995-01-02'}},self.post),'지원자_1995')
  with self.assertRaises(ValueError):r.message('{출생년도}',self.rule,self.post)
 def test_mixed_new_cpa_posting_matches_even_outside_trainee_board(self):
  self.assertTrue(matching.resume_candidate({**self.post,'source':'kicpa:cpa','recruitment_categories':['entry_cpa','experienced_cpa'],'work_types':['full_time'],'career_min_years':2},self.rule['filters']))
