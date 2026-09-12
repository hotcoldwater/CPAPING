"""자유양식 DOCX/PDF와 사용자 대응표가 있는 지정 DOCX를 생성한다."""
import io
import os
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

FIELDS={'full_name':'이름','email':'이메일','phone':'연락처','address':'주소','education':'학력','career':'경력',
        'certifications':'자격·어학','target_firm':'지원 법인','target_role':'지원 직무','essay':'자기소개서'}


def values(profile,email=''):
    data={k:str(profile.get('identity',{}).get(k,'') or '') for k in FIELDS if k!='essay'}
    data['email']=email or data.get('email','')
    data['essay']=str(profile.get('essay','') or '')
    return data


def safe_name(name,extension):
    value=re.sub(r'[^가-힣A-Za-z0-9 ._()-]','_',name or '지원서').strip(' .')[:90]
    value=re.sub(r'\.(pdf|docx)$','',value,flags=re.I) or '지원서'
    return value+'.'+extension


def _docx(raw):
    from docx import Document
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>20_000_000: raise ValueError('양식 압축 해제 용량이 너무 큽니다.')
            if any('vbaProject' in n or n.startswith('word/embeddings/') for n in z.namelist()): raise ValueError('매크로·삽입 개체가 있는 양식은 지원하지 않습니다.')
            # 외부 이미지·연결 필드는 문서 열 때 원격 요청을 유발할 수 있다.
            if any(b'TargetMode="External"' in z.read(n) for n in z.namelist() if n.endswith('.rels')): raise ValueError('외부 연결을 제거한 양식을 올려 주세요.')
        return Document(io.BytesIO(raw))
    except zipfile.BadZipFile as e: raise ValueError('DOCX 파일을 읽을 수 없습니다.') from e


def paragraphs(doc):
    # 본문·표·중첩표·머리말·꼬리말 모두 순회. merged cell은 한 번만.
    seen=set()
    def walk(container):
        for p in container.paragraphs:
            if p._p not in seen: seen.add(p._p); yield p
        for t in container.tables:
            for r in t.rows:
                for c in r.cells: yield from walk(c)
    yield from walk(doc)
    for s in doc.sections:
        for area in (s.header,s.footer,s.first_page_header,s.first_page_footer,s.even_page_header,s.even_page_footer): yield from walk(area)


def inspect(raw,mime):
    if mime=='application/pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if reader.is_encrypted or len(reader.pages)>30: raise ValueError('암호화됐거나 30쪽을 넘는 PDF는 지원하지 않습니다.')
        fields=reader.get_fields() or {}
        if not fields: raise ValueError('입력 필드가 없는 PDF입니다. 편집 가능한 DOCX 양식을 사용해 주세요.')
        return {'type':'pdf','fields':[{'id':k,'label':str(v.get('/TU') or k),'field':k if k in FIELDS else ''} for k,v in fields.items()],
                'help':'각 PDF 입력 필드를 저장한 정보에 연결합니다. 생성 후 한글과 줄 넘침을 반드시 검수해 주세요.'}
    doc=_docx(raw)
    return {'type':'docx','fields':[{'id':str(i),'label':p.text[:240] or '(빈 칸)',
            'field':next((k for k in FIELDS if '{{'+k+'}}' in p.text),'')} for i,p in enumerate(paragraphs(doc))],
            'help':'값을 채울 문단이나 표 칸을 고르세요. 연결한 칸은 값으로 바뀝니다. 제목·안내 문단은 연결하지 마세요.'}


def font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if 'CareerKR' in pdfmetrics.getRegisteredFontNames(): return 'CareerKR'
    candidates=[os.getenv('CAREER_FONT_PATH',''),'/usr/share/fonts/truetype/nanum/NanumGothic.ttf','/Library/Fonts/Arial Unicode.ttf']
    path=next((p for p in candidates if p and Path(p).is_file()),None)
    if not path: raise ValueError('한글 PDF 글꼴 설정이 필요합니다. 운영자에게 문의해 주세요.')
    pdfmetrics.registerFont(TTFont('CareerKR',path)); return 'CareerKR'


def make(profile,format='docx',email='',template=None,mapping=None):
    data=values(profile,email)
    if template:
        raw,mime=template
        if not mapping: raise ValueError('지정 양식의 항목 대응을 저장해 주세요.')
        if any(k not in FIELDS for k in mapping.values()): raise ValueError('양식 항목 대응이 맞지 않습니다.')
        if any(not data.get(k) for k in mapping.values()): raise ValueError('지정 양식에 필요한 기본 정보가 비어 있습니다.')
        if mime=='application/pdf':
            if format!='pdf': raise ValueError('PDF 지정 양식은 PDF로만 제출할 수 있습니다.')
            return fill_pdf(raw,data,mapping)
        if format!='docx': raise ValueError('DOCX 지정 양식은 DOCX로 생성합니다. PDF 요구 공고는 변환·검수가 필요합니다.')
        doc=_docx(raw)
        plist=list(paragraphs(doc));valid={str(i) for i in range(len(plist))}
        if set(mapping)-valid: raise ValueError('양식이 변경되었습니다. 항목 대응을 다시 확인해 주세요.')
        for i,p in enumerate(plist):
            original=p.text
            if str(i) in mapping:
                new=data[mapping[str(i)]]
                # placeholder가 있으면 주위 라벨은 보존한다.
                if '{{'+mapping[str(i)]+'}}' in original: new=original.replace('{{'+mapping[str(i)]+'}}',new)
                if p.runs:
                    p.runs[0].text=new
                    for r in p.runs[1:]: r.text=''
                else: p.add_run(new)
            if '{{' in p.text: raise ValueError('대응하지 않은 양식 항목이 남아 있습니다.')
        out=io.BytesIO();doc.save(out);return out.getvalue()
    if format=='pdf':
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.pagesizes import A4
        name=font();out=io.BytesIO()
        style=ParagraphStyle('body',fontName=name,fontSize=10,leading=17,wordWrap='CJK',spaceAfter=9)
        title=ParagraphStyle('title',parent=style,fontSize=19,leading=28,spaceAfter=18)
        story=[Paragraph('입사지원서',title)]
        for k,label in FIELDS.items():
            if data[k]: story.extend([Paragraph(escape(label),ParagraphStyle(k,parent=style,fontSize=12,spaceBefore=10)),Paragraph(escape(data[k]).replace('\n','<br/>'),style)])
        def page(canvas,doc):
            canvas.setFont(name,8);canvas.drawRightString(A4[0]-48,28,str(doc.page))
        SimpleDocTemplate(out,pagesize=A4,rightMargin=48,leftMargin=48,topMargin=44,bottomMargin=44).build(story,onFirstPage=page,onLaterPages=page)
        return out.getvalue()
    from docx import Document
    from docx.shared import Pt,Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8)
    normal=doc.styles['Normal'];normal.font.name='맑은 고딕';normal.font.size=Pt(10)
    normal.element.get_or_add_rPr().append(OxmlElement('w:rFonts'));normal.element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    doc.add_heading('입사지원서',0)
    for k,label in FIELDS.items():
        if data[k]: doc.add_heading(label,2);doc.add_paragraph(data[k])
    out=io.BytesIO();doc.save(out);return out.getvalue()


def fill_pdf(raw,data,mapping):
    # 기본 PDF 폰트로 한글을 넣으면 깨지므로 폼을 평탄화하고 한글 글꼴로 겹쳐 그린다.
    from pypdf import PdfReader,PdfWriter
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase.pdfmetrics import stringWidth
    reader=PdfReader(io.BytesIO(raw));fields=reader.get_fields() or {}
    if set(mapping)-set(fields): raise ValueError('PDF 필드 대응이 맞지 않습니다.')
    required={k for k,v in fields.items() if int(v.get('/Ff',0)) & 2}
    if required-set(mapping):raise ValueError('PDF의 필수 입력 필드를 모두 연결해 주세요.')
    name=font();writer=PdfWriter();seen=set()
    for original_page in reader.pages:
        page=writer.add_page(original_page)
        if page.rotation or float(page.mediabox.left)!=0 or float(page.mediabox.bottom)!=0: raise ValueError('회전된 PDF 양식은 DOCX로 준비해 주세요.')
        out=io.BytesIO();cv=canvas.Canvas(out,pagesize=(float(page.mediabox.width),float(page.mediabox.height)))
        for ref in page.get('/Annots',[]):
            a=ref.get_object();key=str(a.get('/T') or a.get('/Parent',{}).get_object().get('/T','')) if a.get('/Parent') else str(a.get('/T',''))
            if key not in mapping: continue
            if a.get('/FT',fields[key].get('/FT'))!='/Tx': raise ValueError('텍스트 입력 필드만 자동으로 채울 수 있습니다.')
            x1,y1,x2,y2=map(float,a['/Rect']);width=x2-x1-6;height=y2-y1-6
            text=data[mapping[key]];size=9;leading=12;lines=[]
            for paragraph in text.split('\n'):
                line=''
                for ch in paragraph:
                    if stringWidth(line+ch,name,size)>width:
                        lines.append(line);line=ch
                    else:line+=ch
                lines.append(line)
            if len(lines)*leading>height: raise ValueError(f'PDF 항목 {key}의 공간이 부족합니다. 내용을 줄여 주세요.')
            cv.setFillColorRGB(1,1,1);cv.rect(x1+1,y1+1,x2-x1-2,y2-y1-2,fill=1,stroke=0)
            cv.setFillColorRGB(0,0,0);cv.setFont(name,size)
            for i,line in enumerate(lines):cv.drawString(x1+3,y2-3-size-i*leading,line)
            seen.add(key)
        cv.save();overlay=PdfReader(io.BytesIO(out.getvalue()))
        if overlay.pages: page.merge_page(overlay.pages[0])
        # JavaScript actions·다른 빈 필드가 전달되지 않게 폼 주석을 제거.
        if '/Annots' in page: del page['/Annots']
        if '/AA' in page: del page['/AA']
    if seen!=set(mapping):raise ValueError('PDF의 일부 필드 위치를 읽지 못했습니다.')
    result=io.BytesIO();writer.write(result);return result.getvalue()
